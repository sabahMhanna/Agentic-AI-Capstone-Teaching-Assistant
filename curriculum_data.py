
from pathlib import Path
from docx import Document
from pypdf import PdfReader
import re

PROJECT_DIR = Path(__file__).resolve().parent

COURSE_SPEC_PATH = (
    PROJECT_DIR
    / "knowledge"
    / "AY26-27-CSC51-Course specifications.docx"
)

PYTHON_TUTORIAL_PATH = (
    PROJECT_DIR
    / "knowledge"
    / "PythonTutorial.pdf"
)


# ============================================================
# OFFICIAL SEMESTER SEQUENCE
# ============================================================
# 22 periods = 11 weeks x 2 periods

SEMESTER_PERIODS = [
    "List: collection of data",
    "Sorting list",
    "Sorting list",
    "Processing list",
    "Processing list",
    "Multidimensional array",
    "Multidimensional array",

    "Mutability and sequence",
    "Mutable and immutable data types",
    "Tuple",
    "Tuple",
    "Dictionary",
    "Dictionary",

    "Robotic System: Hardware & Software",
    "MicroPython Programming",
    "AIoT Controllers",
    "AIoT Controllers",
    "Sensors and Data",
    "Sensors and Data",
    "Applications using AI Camera",
    "Applications using AI Camera",
    "Applications using AI Camera",
]


# ============================================================
# WORD TABLE EXTRACTION
# ============================================================

def extract_table_rows(table):
    rows = []

    for row in table.rows:
        values = []

        for cell in row.cells:
            text = cell.text.strip()

            if text:
                values.append(text)

            for nested in cell.tables:
                rows.extend(
                    extract_table_rows(nested)
                )

        if values:
            rows.append(" | ".join(values))

    return rows


def load_course_rows():
    if not COURSE_SPEC_PATH.exists():
        raise FileNotFoundError(
            f"Course specification not found: {COURSE_SPEC_PATH}"
        )

    document = Document(COURSE_SPEC_PATH)

    rows = []

    for table in document.tables:
        rows.extend(
            extract_table_rows(table)
        )

    return rows


# ============================================================
# 11-WEEK MAP
# ============================================================

def build_semester_weeks():
    if len(SEMESTER_PERIODS) != 22:
        raise ValueError(
            "Semester must contain exactly 22 periods."
        )

    return {
        week: {
            "period_1": SEMESTER_PERIODS[(week - 1) * 2],
            "period_2": SEMESTER_PERIODS[(week - 1) * 2 + 1],
        }
        for week in range(1, 12)
    }


def get_week_data(week: int):
    if week < 1 or week > 11:
        raise ValueError(
            "Week must be between 1 and 11."
        )

    return build_semester_weeks()[week]


# ============================================================
# PRECISE PC MAPPING
# ============================================================

def pc_ids_for_topic(topic: str):
    t = topic.strip().lower()

    # Topic 1 - Lists
    if "collection of data" in t:
        return ["PC1.1", "PC1.2"]

    if "sorting" in t:
        return ["PC1.3"]

    if "processing" in t:
        return [
            "PC1.4",
            "PC1.5",
            "PC1.6",
            "PC1.7",
            "PC1.8",
            "PC1.9",
        ]

    if "multidimensional" in t:
        return ["PC1.7", "PC1.10"]

    # Topic 2 - Tuples / dictionaries
    if "mutability and sequence" in t:
        return ["PC2.1", "PC2.2"]

    if "mutable and immutable" in t:
        return ["PC2.1", "PC2.2"]

    if "tuple" in t:
        return [
            "PC2.1",
            "PC2.2",
            "PC2.3",
            "PC2.4",
            "PC2.5",
        ]

    if "dictionary" in t:
        return [
            "PC2.6",
            "PC2.7",
            "PC2.8",
        ]

    # Topic 3 - Robotics / AIoT
    if "robotic system" in t:
        return ["PCA.1"]

    if "micropython" in t:
        return ["PCA.2"]

    if "aiot controllers" in t:
        return ["PCA.3"]

    if "sensors and data" in t:
        return ["PCA.4", "PCA.5"]

    if "ai camera" in t:
        return [
            "PCA.6",
            "PCA.7",
            "PCA.8",
            "PCA.9",
            "PCA.10",
        ]

    return []


def get_performance_criteria(topic: str):
    rows = load_course_rows()
    pc_ids = pc_ids_for_topic(topic)

    if not pc_ids:
        return (
            "No matching performance criteria "
            f"were found for topic: {topic}"
        )

    selected = []

    for row in rows:
        first_field = row.split("|")[0].strip()

        if first_field in pc_ids:
            selected.append(row)

    selected = list(dict.fromkeys(selected))

    if not selected:
        return (
            "No matching performance criteria "
            f"were found for topic: {topic}"
        )

    return "\n".join(selected)


# ============================================================
# TRUSTED RESOURCES
# ============================================================

def get_resources(topic: str):
    t = topic.strip().lower()

    python_essentials_link = (
        "https://www.netacad.com/courses/"
        "python-essentials-1"
        "?courseLang=en-US"
        "&instance_id="
        "c67d18ac-fe42-4e41-a337-09c1dbf540a9"
    )

    if "collection of data" in t:
        resources = [
            "Python Essentials 1 - Cisco Networking Academy",
            f"Course link: {python_essentials_link}",
            "Recommended lesson area: Module 3, Section 3.4 - List",
            "Local knowledge source: PythonTutorial.pdf",
            "Python Videos Folder",
        ]

    elif "sorting" in t:
        resources = [
            "Python Essentials 1 - Cisco Networking Academy",
            f"Course link: {python_essentials_link}",
            "Recommended lesson area: Module 3, Section 3.5 - Sorting Simple List",
            "Local knowledge source: PythonTutorial.pdf",
            "Python Videos Folder",
        ]

    elif "processing" in t:
        resources = [
            "Python Essentials 1 - Cisco Networking Academy",
            f"Course link: {python_essentials_link}",
            "Recommended lesson area: Module 3, Section 3.6 - List Processing",
            "Local knowledge source: PythonTutorial.pdf",
            "Python Videos Folder",
        ]

    elif "multidimensional" in t:
        resources = [
            "Python Essentials 1 - Cisco Networking Academy",
            f"Course link: {python_essentials_link}",
            "Recommended lesson area: Module 3, Section 3.7 - Multidimensional Arrays",
            "Local knowledge source: PythonTutorial.pdf",
            "Python Videos Folder",
        ]

    elif any(
        x in t
        for x in [
            "tuple",
            "mutability",
            "mutable",
            "immutable",
            "dictionary",
        ]
    ):
        resources = [
            "Python Essentials 1 - Cisco Networking Academy",
            f"Course link: {python_essentials_link}",
            "Recommended lesson area: Module 4, Sections 4.1-4.5",
            "Local knowledge source: PythonTutorial.pdf",
            "Python Videos Folder",
        ]

    else:
        resources = [
            "Maker and Coder Portal - MC 4.0 Curriculum",
            "Maker and Coder Portal: https://learningportal.makerandcoder.com/course/us",
            "MC4.0 AIOT Kit",
            "MC Lab Software",
            "Grade 10 Lessons 1-9",
        ]

    return "\n".join(resources)


# ============================================================
# LOCAL PYTHON TUTORIAL PDF RETRIEVAL
# ============================================================

def load_python_tutorial_pages():
    if not PYTHON_TUTORIAL_PATH.exists():
        raise FileNotFoundError(
            f"Python tutorial not found: {PYTHON_TUTORIAL_PATH}"
        )

    reader = PdfReader(
        str(PYTHON_TUTORIAL_PATH)
    )

    pages = []

    for page_number, page in enumerate(
        reader.pages,
        start=1
    ):
        text = page.extract_text() or ""

        text = re.sub(
            r"\s+",
            " ",
            text
        ).strip()

        if text:
            pages.append({
                "page": page_number,
                "text": text,
            })

    return pages


def tutorial_keywords_for_topic(topic: str):
    t = topic.strip().lower()

    if "collection of data" in t:
        return [
            "list",
            "lists",
            "append",
            "insert",
            "index",
            "indices",
            "nested list",
        ]

    if "sorting" in t:
        return [
            "sort",
            "sorted",
            "sorting",
            "reverse",
            "bubble sort",
            "len",
        ]

    if "processing" in t:
        return [
            "slicing",
            "slice",
            "del",
            "in operator",
            "not in",
            "list comprehension",
            "loop",
        ]

    if "multidimensional" in t:
        return [
            "nested list",
            "multidimensional",
            "two dimensional",
            "2d",
            "matrix",
        ]

    if any(
        x in t
        for x in [
            "mutability",
            "mutable",
            "immutable",
            "tuple",
        ]
    ):
        return [
            "tuple",
            "immutable",
            "sequence",
            "mutable",
        ]

    if "dictionary" in t:
        return [
            "dictionary",
            "dict",
            "key",
            "value",
            "items",
        ]

    return []


def get_tutorial_material(
    topic: str,
    max_pages: int = 4,
    max_chars_per_page: int = 2600
):
    # Retrieve the most relevant text pages from PythonTutorial.pdf.
    # PAGE markers are preserved so the Resource Agent can cite the
    # source and Step 12 can render that PDF page as a slide visual.

    keywords = tutorial_keywords_for_topic(topic)

    if not keywords:
        return (
            "No PythonTutorial.pdf retrieval is configured "
            f"for this non-Python topic: {topic}"
        )

    pages = load_python_tutorial_pages()

    scored = []

    for item in pages:
        lower_text = item["text"].lower()

        score = 0

        for keyword in keywords:
            keyword_lower = keyword.lower()

            # Phrase hits get a little extra weight.
            hits = lower_text.count(keyword_lower)
            score += hits * (
                3 if " " in keyword_lower else 1
            )

        if score > 0:
            scored.append({
                "score": score,
                "page": item["page"],
                "text": item["text"],
            })

    scored.sort(
        key=lambda x: (
            -x["score"],
            x["page"]
        )
    )

    selected = scored[:max_pages]

    if not selected:
        return (
            "No relevant material was found in "
            f"PythonTutorial.pdf for topic: {topic}"
        )

    chunks = []

    for item in selected:
        excerpt = item["text"][:max_chars_per_page]

        chunks.append(
            "SOURCE: PythonTutorial.pdf\n"
            f"PAGE: {item['page']}\n"
            f"RELEVANCE SCORE: {item['score']}\n"
            f"CONTENT:\n{excerpt}"
        )

    return "\n\n---\n\n".join(chunks)


def get_tutorial_pages_for_topic(topic: str):
    # Return only the best matching page numbers.
    # Useful for diagnostics and PowerPoint source visuals.

    material = get_tutorial_material(topic)

    return [
        int(x)
        for x in re.findall(
            r"PAGE:\s*(\d+)",
            material
        )
    ]
